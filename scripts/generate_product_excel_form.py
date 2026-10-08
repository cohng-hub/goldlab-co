# -*- coding: utf-8 -*-
"""
Generate a luxury, professional Gold Lab Product Registration & Inventory Excel Workbook.
Includes:
1. Product Detail Registration Card (제품상세_등록카드) with dedicated Photo Registration Box
2. Product Catalog & Inventory Sheet (제품_목록관리대장) with Row Photo Thumbnails & KPI Dashboard
3. User Guide & Settings Sheet (사용안내_및_설정) with dropdown validations & photo insertion tips
"""

import os
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.drawing.image import Image as OpenpyxlImage
from PIL import Image as PILImage, ImageDraw, ImageFont

# Base directory
BASE_DIR = r"c:\Users\WD\Desktop\작업용\ju\골드랩 작업"
SCRATCH_DIR = os.path.join(BASE_DIR, "scratch", "excel_assets")
os.makedirs(SCRATCH_DIR, exist_ok=True)

# -------------------------------------------------------------
# 1. GENERATE SAMPLE IMAGES FOR THE EXCEL FORM
# -------------------------------------------------------------
def create_sample_placeholder():
    path = os.path.join(SCRATCH_DIR, "photo_placeholder.png")
    img = PILImage.new("RGBA", (340, 300), (248, 250, 252, 255))
    draw = ImageDraw.Draw(img)
    
    # Outer gold border
    draw.rectangle([6, 6, 333, 293], outline=(212, 175, 55, 255), width=2)
    # Inner dashed line effect
    for x in range(16, 324, 16):
        draw.line([(x, 16), (x + 8, 16)], fill=(203, 213, 225, 255), width=2)
        draw.line([(x, 283), (x + 8, 283)], fill=(203, 213, 225, 255), width=2)
    for y in range(16, 284, 16):
        draw.line([(16, y), (16, y + 8)], fill=(203, 213, 225, 255), width=2)
        draw.line([(323, y), (323, y + 8)], fill=(203, 213, 225, 255), width=2)
        
    # Camera Icon
    cx, cy = 170, 115
    draw.rounded_rectangle([cx - 45, cy - 25, cx + 45, cy + 35], radius=8, fill=(241, 245, 249, 255), outline=(148, 163, 184, 255), width=2)
    draw.rounded_rectangle([cx - 20, cy - 36, cx + 20, cy - 24], radius=3, fill=(203, 213, 225, 255), outline=(148, 163, 184, 255), width=2)
    draw.ellipse([cx - 20, cy - 10, cx + 20, cy + 30], fill=(212, 175, 55, 255), outline=(184, 134, 11, 255), width=2)
    draw.ellipse([cx - 10, cy, cx + 10, cy + 20], fill=(30, 41, 59, 255))
    
    try:
        font_main = ImageFont.truetype("malgun.ttf", 15)
        font_sub = ImageFont.truetype("malgun.ttf", 12)
    except:
        font_main = font_sub = ImageFont.load_default()
        
    draw.text((170, 185), "제품 실물 사진 등록 영역", fill=(30, 41, 59, 255), font=font_main, anchor="mm")
    draw.text((170, 215), "[메뉴] 삽입 > 그림 > 이 디바이스", fill=(100, 116, 139, 255), font=font_sub, anchor="mm")
    draw.text((170, 238), "(또는 최신 엑셀 '셀에 그림 삽입')", fill=(184, 134, 11, 255), font=font_sub, anchor="mm")
    
    img.save(path, "PNG")
    return path

def create_sample_goldbar():
    path = os.path.join(SCRATCH_DIR, "sample_goldbar.png")
    img = PILImage.new("RGBA", (110, 80), (255, 255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.rounded_rectangle([4, 4, 105, 75], radius=6, fill=(254, 243, 199, 255), outline=(212, 175, 55, 255), width=2)
    draw.rounded_rectangle([18, 15, 92, 65], radius=4, fill=(245, 158, 11, 255), outline=(180, 83, 9, 255), width=1)
    draw.line([(22, 19), (88, 19)], fill=(253, 230, 138, 255), width=2)
    try:
        f = ImageFont.truetype("malgun.ttf", 11)
        f_sub = ImageFont.truetype("malgun.ttf", 9)
    except:
        f = f_sub = ImageFont.load_default()
    draw.text((55, 30), "순금 24K", fill=(255, 255, 255, 255), font=f, anchor="mm")
    draw.text((55, 48), "10돈 (37.5g)", fill=(254, 243, 199, 255), font=f_sub, anchor="mm")
    img.save(path, "PNG")
    return path

def create_sample_necklace():
    path = os.path.join(SCRATCH_DIR, "sample_necklace.png")
    img = PILImage.new("RGBA", (110, 80), (255, 255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.rounded_rectangle([4, 4, 105, 75], radius=6, fill=(248, 250, 252, 255), outline=(203, 213, 225, 255), width=2)
    draw.arc([25, 12, 85, 55], start=0, end=180, fill=(212, 175, 55, 255), width=3)
    draw.polygon([(55, 38), (63, 48), (55, 58), (47, 48)], fill=(212, 175, 55, 255), outline=(180, 83, 9, 255))
    draw.ellipse([51, 44, 59, 52], fill=(255, 255, 255, 255))
    try:
        f = ImageFont.truetype("malgun.ttf", 10)
    except:
        f = ImageFont.load_default()
    draw.text((55, 68), "18K 목걸이", fill=(30, 41, 59, 255), font=f, anchor="mm")
    img.save(path, "PNG")
    return path

def create_sample_ring():
    path = os.path.join(SCRATCH_DIR, "sample_ring.png")
    img = PILImage.new("RGBA", (110, 80), (255, 255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.rounded_rectangle([4, 4, 105, 75], radius=6, fill=(248, 250, 252, 255), outline=(203, 213, 225, 255), width=2)
    draw.ellipse([35, 26, 75, 66], outline=(203, 213, 225, 255), width=4)
    draw.polygon([(55, 14), (64, 25), (55, 30), (46, 25)], fill=(224, 242, 254, 255), outline=(56, 189, 248, 255), width=1)
    try:
        f = ImageFont.truetype("malgun.ttf", 10)
    except:
        f = ImageFont.load_default()
    draw.text((55, 68), "14K 다이아반지", fill=(30, 41, 59, 255), font=f, anchor="mm")
    img.save(path, "PNG")
    return path

placeholder_path = create_sample_placeholder()
goldbar_path = create_sample_goldbar()
necklace_path = create_sample_necklace()
ring_path = create_sample_ring()
print("Sample images generated successfully.")

# -------------------------------------------------------------
# 2. CREATE WORKBOOK & DEFINE COLOR STYLES
# -------------------------------------------------------------
wb = openpyxl.Workbook()

# Style tokens
FONT_NAME = "맑은 고딕"
COLOR_NAVY = "1E293B"      # Deep luxury navy
COLOR_GOLD = "D4AF37"      # Gold accent
COLOR_GOLD_DARK = "B8860B" # Dark gold
COLOR_GOLD_BG = "FFFBEB"   # Soft gold tint for formulas
COLOR_HEADER_BG = "0F172A" # Darker header
COLOR_ROW_ALT = "F8FAFC"   # Alternating row
COLOR_BORDER = "CBD5E1"    # Soft slate border
COLOR_BORDER_GOLD = "D4AF37"

font_title = Font(name=FONT_NAME, size=16, bold=True, color="FFFFFF")
font_subtitle = Font(name=FONT_NAME, size=9, bold=False, color="94A3B8")
font_section = Font(name=FONT_NAME, size=11, bold=True, color="1E293B")
font_th = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
font_cell = Font(name=FONT_NAME, size=10, bold=False, color="1E293B")
font_cell_bold = Font(name=FONT_NAME, size=10, bold=True, color="1E293B")
font_formula = Font(name=FONT_NAME, size=10, bold=True, color="1E40AF")
font_kpi_num = Font(name=FONT_NAME, size=14, bold=True, color="1E293B")

fill_navy_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
fill_section_header = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
fill_label = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
fill_formula = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
fill_alt = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
fill_card_bg = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
fill_kpi_bg = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

thin_border_side = Side(border_style="thin", color=COLOR_BORDER)
border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")

# -------------------------------------------------------------
# 3. SHEET 1: 제품상세_등록카드 (Single Product Detail Card)
# -------------------------------------------------------------
ws1 = wb.active
ws1.title = "제품상세_등록카드"
ws1.views.sheetView[0].showGridLines = True

# Column widths
ws1.column_dimensions['A'].width = 3
ws1.column_dimensions['B'].width = 13
ws1.column_dimensions['C'].width = 14
ws1.column_dimensions['D'].width = 14
ws1.column_dimensions['E'].width = 3
ws1.column_dimensions['F'].width = 14
ws1.column_dimensions['G'].width = 22
ws1.column_dimensions['H'].width = 14
ws1.column_dimensions['I'].width = 22
ws1.column_dimensions['J'].width = 3

# Top Header Banner (Rows 2 to 3, Cols B to I)
ws1.merge_cells("B2:I2")
ws1["B2"] = "💎 골드랩(GOLD LAB) 프리미엄 제품 상세 등록 카드"
ws1["B2"].font = font_title
ws1["B2"].fill = fill_navy_header
ws1["B2"].alignment = Alignment(horizontal="center", vertical="center")
ws1.row_dimensions[2].height = 40

ws1.merge_cells("B3:I3")
ws1["B3"] = "등록일시: 2026-10-06  |  관리번호: GL-2026-001  |  작성자: 골드랩 관리자  |  * [수식(연노랑)] 표기 셀은 자동 계산됩니다."
ws1["B3"].font = font_subtitle
ws1["B3"].fill = fill_navy_header
ws1["B3"].alignment = Alignment(horizontal="center", vertical="center")
ws1.row_dimensions[3].height = 20

# Row 4 blank
ws1.row_dimensions[4].height = 10

# Left Section: Photo Box (B5:D20)
ws1.merge_cells("B5:D5")
ws1["B5"] = "📷 제품 실물 사진"
ws1["B5"].font = font_section
ws1["B5"].fill = fill_section_header
ws1["B5"].alignment = align_center
ws1["B5"].border = border_cell
ws1.row_dimensions[5].height = 24

ws1.merge_cells("B6:D19")
for r in range(6, 20):
    ws1.row_dimensions[r].height = 18
    for c in range(2, 5):
        cell = ws1.cell(row=r, column=c)
        cell.border = border_cell
        cell.fill = PatternFill(start_color="FAFAFA", end_color="FAFAFA", fill_type="solid")

img_placeholder = OpenpyxlImage(placeholder_path)
img_placeholder.width = 260
img_placeholder.height = 230
ws1.add_image(img_placeholder, "B6")

# Photo Guide Note under photo box
ws1.merge_cells("B20:D21")
ws1["B20"] = "💡 사진 등록 팁: 상단 [삽입] > [그림] > [이 디바이스]로 사진 삽입 후, 우클릭 > [그림 서식] > [크기 및 속성] > '위치와 크기 변동'을 체크하시면 완벽하게 고정됩니다."
ws1["B20"].font = Font(name=FONT_NAME, size=8, color="64748B")
ws1["B20"].fill = fill_label
ws1["B20"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws1["B20"].border = border_cell

# Right Section: Basic & Spec Info (F5:I14)
ws1.merge_cells("F5:I5")
ws1["F5"] = "📋 제품 기본 정보 및 규격 사양"
ws1["F5"].font = font_section
ws1["F5"].fill = fill_section_header
ws1["F5"].alignment = align_left
ws1["F5"].border = border_cell

spec_fields = [
    (6, "상품 관리코드", "GL-2026-001", "카테고리/분류", "순금 골드바"),
    (7, "제 품 명", "24K 한국표준 순금 골드바 10돈", "거래 구분", "신품 입고"),
    (8, "재질 및 순도", "24K 순금 (99.99%)", "각인/홀마크", "Hallmarked 999.9"),
    (9, "실측 중량(g)", 37.50, "환산 중량(돈)", "=IF(ISNUMBER(G9), ROUND(G9/3.75, 2), 0)"),
    (10, "제품 크기/치수", "가로 22mm x 세로 36mm", "보증서 유무", "품질보증서 동봉"),
    (11, "제조/공급처", "한국표준금거래소 종로본사", "입고 일자", "2026-10-06"),
    (12, "시리얼 번호", "SN-9948201", "담당 감정사", "김골드 공인감정원"),
    (13, "원산지", "대한민국 (국내제조)", "패키지 상태", "보안 밀봉 블리스터팩")
]

for row_idx, l1, v1, l2, v2 in spec_fields:
    ws1.row_dimensions[row_idx].height = 24
    # Col F
    c_f = ws1.cell(row=row_idx, column=6, value=l1)
    c_f.font = font_cell_bold
    c_f.fill = fill_label
    c_f.alignment = align_center
    c_f.border = border_cell
    # Col G
    c_g = ws1.cell(row=row_idx, column=7, value=v1)
    c_g.font = font_cell
    c_g.alignment = align_center if row_idx != 7 else align_left
    c_g.border = border_cell
    if row_idx == 9:
        c_g.number_format = '0.00 "g"'
    # Col H
    c_h = ws1.cell(row=row_idx, column=8, value=l2)
    c_h.font = font_cell_bold
    c_h.fill = fill_label
    c_h.alignment = align_center
    c_h.border = border_cell
    # Col I
    c_i = ws1.cell(row=row_idx, column=9, value=v2)
    c_i.border = border_cell
    if row_idx == 9: # Formula for Don
        c_i.font = font_formula
        c_i.fill = fill_formula
        c_i.alignment = align_center
        c_i.number_format = '0.00 "돈"'
    else:
        c_i.font = font_cell
        c_i.alignment = align_center

# Right Section: Pricing & Margin (F15:I19)
ws1.merge_cells("F15:I15")
ws1["F15"] = "💰 매입 원가 및 판매가 산정 (자동 계산)"
ws1["F15"].font = font_section
ws1["F15"].fill = fill_section_header
ws1["F15"].alignment = align_left
ws1["F15"].border = border_cell
ws1.row_dimensions[15].height = 24

price_fields = [
    (16, "매입 단가(원)", 6800000, "가공/공임비(원)", 30000),
    (17, "총 매입원가", "=G16+I16", "부가세(VAT)", "=ROUND(G17*0.1, 0)"),
    (18, "판매 희망가(원)", 7350000, "예상 마진액(원)", "=G18-G17"),
    (19, "예상 마진율", "=IF(G17>0, I18/G17, 0)", "결제/거래방식", "세금계산서 / 현금영수증")
]

for row_idx, l1, v1, l2, v2 in price_fields:
    ws1.row_dimensions[row_idx].height = 24
    # Col F
    c_f = ws1.cell(row=row_idx, column=6, value=l1)
    c_f.font = font_cell_bold
    c_f.fill = fill_label
    c_f.alignment = align_center
    c_f.border = border_cell
    # Col G
    c_g = ws1.cell(row=row_idx, column=7, value=v1)
    c_g.border = border_cell
    if row_idx in (16, 18):
        c_g.font = font_cell_bold
        c_g.alignment = align_right
        c_g.number_format = '#,##0 "원"'
    elif row_idx == 17:
        c_g.font = font_formula
        c_g.fill = fill_formula
        c_g.alignment = align_right
        c_g.number_format = '#,##0 "원"'
    elif row_idx == 19:
        c_g.font = font_formula
        c_g.fill = fill_formula
        c_g.alignment = align_center
        c_g.number_format = '0.0%'
        
    # Col H
    c_h = ws1.cell(row=row_idx, column=8, value=l2)
    c_h.font = font_cell_bold
    c_h.fill = fill_label
    c_h.alignment = align_center
    c_h.border = border_cell
    # Col I
    c_i = ws1.cell(row=row_idx, column=9, value=v2)
    c_i.border = border_cell
    if row_idx == 16:
        c_i.font = font_cell
        c_i.alignment = align_right
        c_i.number_format = '#,##0 "원"'
    elif row_idx in (17, 18):
        c_i.font = font_formula
        c_i.fill = fill_formula
        c_i.alignment = align_right
        c_i.number_format = '#,##0 "원"'
    elif row_idx == 19:
        c_i.font = font_cell
        c_i.alignment = align_center

# Bottom Section 1: Inventory & Storage (Rows 23 to 26)
ws1.merge_cells("B23:I23")
ws1["B23"] = "📦 재고 보유 및 보관 위치 관리"
ws1["B23"].font = font_section
ws1["B23"].fill = fill_section_header
ws1["B23"].alignment = align_left
ws1["B23"].border = border_cell
ws1.row_dimensions[23].height = 24

inv_fields = [
    (24, "보유 재고수량", 2, "보관 위치", "본점 쇼케이스 A-01 (골드바 전용)"),
    (25, "판매/진열 상태", "진열 판매중", "재고 평가총액", "=G17*C24")
]

for row_idx, l1, v1, l2, v2 in inv_fields:
    ws1.row_dimensions[row_idx].height = 24
    c_b = ws1.cell(row=row_idx, column=2, value=l1)
    c_b.font = font_cell_bold
    c_b.fill = fill_label
    c_b.alignment = align_center
    c_b.border = border_cell
    
    ws1.merge_cells(start_row=row_idx, start_column=3, end_row=row_idx, end_column=4)
    c_c = ws1.cell(row=row_idx, column=3, value=v1)
    c_c.font = font_cell_bold
    c_c.alignment = align_center
    for c in range(3, 5):
        ws1.cell(row=row_idx, column=c).border = border_cell
    if row_idx == 24:
        c_c.number_format = '#,##0 "개"'
        
    c_f = ws1.cell(row=row_idx, column=6, value=l2)
    c_f.font = font_cell_bold
    c_f.fill = fill_label
    c_f.alignment = align_center
    c_f.border = border_cell
    
    ws1.merge_cells(start_row=row_idx, start_column=7, end_row=row_idx, end_column=9)
    c_g = ws1.cell(row=row_idx, column=7, value=v2)
    for c in range(7, 10):
        ws1.cell(row=row_idx, column=c).border = border_cell
    if row_idx == 25:
        c_g.font = font_formula
        c_g.fill = fill_formula
        c_g.alignment = align_right
        c_g.number_format = '#,##0 "원"'
    else:
        c_g.font = font_cell
        c_g.alignment = align_left

# Bottom Section 2: Remarks / Appraisal Memo (Rows 27 to 31)
ws1.merge_cells("B27:I27")
ws1["B27"] = "📝 제품 특이사항 및 감정/고객 안내 메모"
ws1["B27"].font = font_section
ws1["B27"].fill = fill_section_header
ws1["B27"].alignment = align_left
ws1["B27"].border = border_cell
ws1.row_dimensions[27].height = 24

ws1.merge_cells("B28:I31")
memo_text = (
    "• 순금 999.9 포장 밀봉 블리스터 팩 상태 최상 (미개봉 신품)\n"
    "• 한국표준금거래소 정품 각인 및 시리얼 넘버 보증서 일치 확인 완료\n"
    "• 고객 구매 시 전용 고급 가죽 케이스 및 정품 보증서 동봉 출고 요망\n"
    "• 당일 종로 실시간 금 매입/판매 기준가 변동에 따라 단가 실시간 확인 필요"
)
ws1["B28"] = memo_text
ws1["B28"].font = font_cell
ws1["B28"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
for r in range(28, 32):
    ws1.row_dimensions[r].height = 20
    for c in range(2, 10):
        ws1.cell(row=r, column=c).border = border_cell
        ws1.cell(row=r, column=c).fill = fill_card_bg

# -------------------------------------------------------------
# 4. SHEET 2: 제품_목록관리대장 (Catalog Table & Inventory Sheet)
# -------------------------------------------------------------
ws2 = wb.create_sheet(title="제품_목록관리대장")
ws2.views.sheetView[0].showGridLines = True

# Column widths
cols_config = [
    ('A', 3),
    ('B', 6),   # No.
    ('C', 18),  # 제품 사진 (Photo Thumbnail)
    ('D', 15),  # 상품코드
    ('E', 14),  # 카테고리
    ('F', 24),  # 상품명
    ('G', 15),  # 순도
    ('H', 12),  # 중량(g)
    ('I', 12),  # 중량(돈) - 수식
    ('J', 13),  # 입고일자
    ('K', 14),  # 매입단가
    ('L', 12),  # 공임비
    ('M', 15),  # 총 매입원가 - 수식
    ('N', 15),  # 판매단가
    ('O', 11),  # 재고수량
    ('P', 16),  # 재고평가총액 - 수식
    ('Q', 16),  # 예상판매총액 - 수식
    ('R', 12),  # 마진율(%) - 수식
    ('S', 12),  # 판매상태
    ('T', 15),  # 보관위치
    ('U', 25)   # 비고 및 보증서번호
]
for col_letter, width in cols_config:
    ws2.column_dimensions[col_letter].width = width

# Top Title Banner
ws2.merge_cells("B2:U2")
ws2["B2"] = "📊 골드랩(GOLD LAB) 제품 목록 및 재고 관리 대장"
ws2["B2"].font = font_title
ws2["B2"].fill = fill_navy_header
ws2["B2"].alignment = Alignment(horizontal="center", vertical="center")
ws2.row_dimensions[2].height = 36

# Top Dashboard KPI Summary Cards (Rows 3 to 4)
# Card 1: 총 품목수 (B3:D4)
ws2.merge_cells("B3:D3")
ws2["B3"] = "총 등록 품목"
ws2["B3"].font = font_subtitle
ws2["B3"].fill = fill_section_header
ws2["B3"].alignment = align_center

ws2.merge_cells("B4:D4")
ws2["B4"] = '=COUNTA(D8:D30)'
ws2["B4"].font = font_kpi_num
ws2["B4"].fill = fill_kpi_bg
ws2["B4"].alignment = align_center
ws2["B4"].number_format = '#,##0 "개 품목"'

# Card 2: 총 재고수량 (E3:G4)
ws2.merge_cells("E3:G3")
ws2["E3"] = "총 보유 재고수량"
ws2["E3"].font = font_subtitle
ws2["E3"].fill = fill_section_header
ws2["E3"].alignment = align_center

ws2.merge_cells("E4:G4")
ws2["E4"] = '=SUM(O8:O30)'
ws2["E4"].font = font_kpi_num
ws2["E4"].fill = fill_kpi_bg
ws2["E4"].alignment = align_center
ws2["E4"].number_format = '#,##0 "개"'

# Card 3: 총 재고 원가 평가액 (H3:L4)
ws2.merge_cells("H3:L3")
ws2["H3"] = "총 재고 원가 평가액"
ws2["H3"].font = font_subtitle
ws2["H3"].fill = fill_section_header
ws2["H3"].alignment = align_center

ws2.merge_cells("H4:L4")
ws2["H4"] = '=SUM(P8:P30)'
ws2["H4"].font = Font(name=FONT_NAME, size=14, bold=True, color="B8860B")
ws2["H4"].fill = fill_kpi_bg
ws2["H4"].alignment = align_center
ws2["H4"].number_format = '#,##0 "원"'

# Card 4: 예상 총 판매금액 (M3:Q4)
ws2.merge_cells("M3:Q3")
ws2["M3"] = "예상 총 판매 매출액"
ws2["M3"].font = font_subtitle
ws2["M3"].fill = fill_section_header
ws2["M3"].alignment = align_center

ws2.merge_cells("M4:Q4")
ws2["M4"] = '=SUM(Q8:Q30)'
ws2["M4"].font = Font(name=FONT_NAME, size=14, bold=True, color="1E40AF")
ws2["M4"].fill = fill_kpi_bg
ws2["M4"].alignment = align_center
ws2["M4"].number_format = '#,##0 "원"'

# Card 5: 평균 마진율 (R3:U4)
ws2.merge_cells("R3:U3")
ws2["R3"] = "평균 예상 마진율"
ws2["R3"].font = font_subtitle
ws2["R3"].fill = fill_section_header
ws2["R3"].alignment = align_center

ws2.merge_cells("R4:U4")
ws2["R4"] = '=IFERROR(AVERAGE(R8:R30), 0)'
ws2["R4"].font = Font(name=FONT_NAME, size=14, bold=True, color="059669")
ws2["R4"].fill = fill_kpi_bg
ws2["R4"].alignment = align_center
ws2["R4"].number_format = '0.0%'

for r in range(3, 5):
    for c in range(2, 22):
        ws2.cell(row=r, column=c).border = border_cell

# Row 5-6 blank / guide note
ws2.merge_cells("B6:U6")
ws2["B6"] = "💡 팁: '제품 사진' 칸(C열)에 사진을 넣으실 때 최신 엑셀은 [삽입 > 그림 > 셀에 배치], 일반 엑셀은 [이 디바이스]로 넣으신 후 우클릭 > [그림 서식] > '위치와 크기 변동'을 설정하세요."
ws2["B6"].font = Font(name=FONT_NAME, size=8.5, color="64748B")
ws2["B6"].alignment = align_left
ws2.row_dimensions[6].height = 18

# Table Header Row (Row 7)
headers = [
    ("No.", align_center),
    ("제품 사진\n(미리보기)", align_center),
    ("상품코드", align_center),
    ("카테고리", align_center),
    ("상품명", align_center),
    ("순도/재질", align_center),
    ("중량(g)", align_center),
    ("중량(돈)\n[자동환산]", align_center),
    ("입고일자", align_center),
    ("매입단가(원)", align_center),
    ("공임비(원)", align_center),
    ("총 매입원가\n[자동합계]", align_center),
    ("판매단가(원)", align_center),
    ("재고\n수량", align_center),
    ("재고평가총액\n[자동계산]", align_center),
    ("예상판매총액\n[자동계산]", align_center),
    ("마진율(%)\n[자동계산]", align_center),
    ("판매상태", align_center),
    ("보관위치", align_center),
    ("비고 및 보증서번호", align_center)
]

ws2.row_dimensions[7].height = 32
for col_idx, (h_title, h_align) in enumerate(headers, start=2):
    cell = ws2.cell(row=7, column=col_idx, value=h_title)
    cell.font = font_th
    cell.fill = fill_navy_header
    cell.alignment = h_align
    cell.border = border_cell

# Sample Product Rows (Rows 8 to 10)
sample_rows = [
    (
        1, "GL-2026-001", "순금(24K)", "24K 한국표준 순금 골드바 10돈", "24K (99.99%)",
        37.50, "2026-10-06", 6800000, 30000, 7350000, 2, "진열 판매중", "쇼케이스 A-01", "정품 보증서 동봉 / 밀봉팩"
    ),
    (
        2, "GL-2026-002", "18K 골드", "18K 반클리프 스타일 클로버 목걸이", "18K (75.0%)",
        5.62, "2026-10-05", 820000, 95000, 1150000, 3, "진열 판매중", "쇼케이스 B-03", "자개 펜던트 세팅 / 선물포장"
    ),
    (
        3, "GL-2026-003", "다이아/보석", "14K 화이트골드 3부 다이아 솔리테어 링", "14K (58.5%)",
        3.15, "2026-10-04", 1250000, 150000, 1790000, 1, "예약중", "금고 1호실", "우신 0.3ct F SI1 VG 감정서 동봉"
    )
]

for idx, data in enumerate(sample_rows):
    row_num = 8 + idx
    ws2.row_dimensions[row_num].height = 65
    no, code, cat, name, purity, gram, in_date, cost, labor, price, stock, status, loc, note = data
    
    # B: No
    c_b = ws2.cell(row=row_num, column=2, value=no)
    c_b.alignment = align_center
    c_b.font = font_cell
    
    # C: Photo cell
    c_c = ws2.cell(row=row_num, column=3)
    c_c.alignment = align_center
    
    # D: Code
    c_d = ws2.cell(row=row_num, column=4, value=code)
    c_d.alignment = align_center
    c_d.font = font_cell_bold
    
    # E: Category
    c_e = ws2.cell(row=row_num, column=5, value=cat)
    c_e.alignment = align_center
    c_e.font = font_cell
    
    # F: Product Name
    c_f = ws2.cell(row=row_num, column=6, value=name)
    c_f.alignment = align_left
    c_f.font = font_cell_bold
    
    # G: Purity
    c_g = ws2.cell(row=row_num, column=7, value=purity)
    c_g.alignment = align_center
    c_g.font = font_cell
    
    # H: Weight (g)
    c_h = ws2.cell(row=row_num, column=8, value=gram)
    c_h.alignment = align_right
    c_h.font = font_cell
    c_h.number_format = '0.00'
    
    # I: Weight (Don) - Robust Formula
    c_i = ws2.cell(row=row_num, column=9, value=f'=IF(ISNUMBER(H{row_num}), ROUND(H{row_num}/3.75, 2), "")')
    c_i.alignment = align_center
    c_i.font = font_formula
    c_i.fill = fill_formula
    c_i.number_format = '0.00 "돈"'
    
    # J: Date
    c_j = ws2.cell(row=row_num, column=10, value=in_date)
    c_j.alignment = align_center
    c_j.font = font_cell
    
    # K: Purchase Cost
    c_k = ws2.cell(row=row_num, column=11, value=cost)
    c_k.alignment = align_right
    c_k.font = font_cell
    c_k.number_format = '#,##0'
    
    # L: Labor Fee
    c_l = ws2.cell(row=row_num, column=12, value=labor)
    c_l.alignment = align_right
    c_l.font = font_cell
    c_l.number_format = '#,##0'
    
    # M: Total Cost (K + L) - Robust Formula
    c_m = ws2.cell(row=row_num, column=13, value=f'=IF(ISNUMBER(K{row_num}), K{row_num}+IF(ISNUMBER(L{row_num}), L{row_num}, 0), "")')
    c_m.alignment = align_right
    c_m.font = font_formula
    c_m.fill = fill_formula
    c_m.number_format = '#,##0'
    
    # N: Selling Price
    c_n = ws2.cell(row=row_num, column=14, value=price)
    c_n.alignment = align_right
    c_n.font = font_cell_bold
    c_n.number_format = '#,##0'
    
    # O: Stock Qty
    c_o = ws2.cell(row=row_num, column=15, value=stock)
    c_o.alignment = align_center
    c_o.font = font_cell_bold
    c_o.number_format = '#,##0'
    
    # P: Inventory Total Value (M * O) - Robust Formula
    c_p = ws2.cell(row=row_num, column=16, value=f'=IF(AND(ISNUMBER(M{row_num}), ISNUMBER(O{row_num})), M{row_num}*O{row_num}, "")')
    c_p.alignment = align_right
    c_p.font = font_formula
    c_p.fill = fill_formula
    c_p.number_format = '#,##0'
    
    # Q: Expected Total Sales (N * O) - Robust Formula
    c_q = ws2.cell(row=row_num, column=17, value=f'=IF(AND(ISNUMBER(N{row_num}), ISNUMBER(O{row_num})), N{row_num}*O{row_num}, "")')
    c_q.alignment = align_right
    c_q.font = font_formula
    c_q.fill = fill_formula
    c_q.number_format = '#,##0'
    
    # R: Margin Rate ((N - M) / M) - Robust Formula
    c_r = ws2.cell(row=row_num, column=18, value=f'=IF(AND(ISNUMBER(M{row_num}), M{row_num}>0, ISNUMBER(N{row_num})), (N{row_num}-M{row_num})/M{row_num}, "")')
    c_r.alignment = align_center
    c_r.font = font_formula
    c_r.fill = fill_formula
    c_r.number_format = '0.0%'
    
    # S: Status
    c_s = ws2.cell(row=row_num, column=19, value=status)
    c_s.alignment = align_center
    c_s.font = font_cell_bold
    
    # T: Location
    c_t = ws2.cell(row=row_num, column=20, value=loc)
    c_t.alignment = align_center
    c_t.font = font_cell
    
    # U: Note
    c_u = ws2.cell(row=row_num, column=21, value=note)
    c_u.alignment = align_left
    c_u.font = font_cell

    for c in range(2, 22):
        ws2.cell(row=row_num, column=c).border = border_cell

# Add Sample Thumbnails to Sheet 2
thumb_goldbar = OpenpyxlImage(goldbar_path)
thumb_goldbar.width = 100
thumb_goldbar.height = 70
ws2.add_image(thumb_goldbar, "C8")

thumb_necklace = OpenpyxlImage(necklace_path)
thumb_necklace.width = 100
thumb_necklace.height = 70
ws2.add_image(thumb_necklace, "C9")

thumb_ring = OpenpyxlImage(ring_path)
thumb_ring.width = 100
thumb_ring.height = 70
ws2.add_image(thumb_ring, "C10")

# Pre-populate empty rows (11 to 30) with formulas & styling ready to type
for row_num in range(11, 31):
    ws2.row_dimensions[row_num].height = 65
    is_alt = (row_num % 2 == 1)
    
    # B: No
    c_b = ws2.cell(row=row_num, column=2, value=row_num - 7)
    c_b.alignment = align_center
    c_b.font = font_cell
    
    # C: Photo cell
    c_c = ws2.cell(row=row_num, column=3)
    c_c.alignment = align_center
    
    # D: Code
    c_d = ws2.cell(row=row_num, column=4)
    c_d.alignment = align_center
    c_d.font = font_cell_bold
    
    # E: Category
    c_e = ws2.cell(row=row_num, column=5)
    c_e.alignment = align_center
    c_e.font = font_cell
    
    # F: Product Name
    c_f = ws2.cell(row=row_num, column=6)
    c_f.alignment = align_left
    c_f.font = font_cell_bold
    
    # G: Purity
    c_g = ws2.cell(row=row_num, column=7)
    c_g.alignment = align_center
    c_g.font = font_cell
    
    # H: Weight (g)
    c_h = ws2.cell(row=row_num, column=8)
    c_h.alignment = align_right
    c_h.font = font_cell
    c_h.number_format = '0.00'
    
    # I: Weight (Don) - Robust Formula
    c_i = ws2.cell(row=row_num, column=9, value=f'=IF(ISNUMBER(H{row_num}), ROUND(H{row_num}/3.75, 2), "")')
    c_i.alignment = align_center
    c_i.font = font_formula
    c_i.fill = fill_formula
    c_i.number_format = '0.00 "돈"'
    
    # J: Date
    c_j = ws2.cell(row=row_num, column=10)
    c_j.alignment = align_center
    c_j.font = font_cell
    
    # K: Purchase Cost
    c_k = ws2.cell(row=row_num, column=11)
    c_k.alignment = align_right
    c_k.font = font_cell
    c_k.number_format = '#,##0'
    
    # L: Labor Fee
    c_l = ws2.cell(row=row_num, column=12)
    c_l.alignment = align_right
    c_l.font = font_cell
    c_l.number_format = '#,##0'
    
    # M: Total Cost (K + L) - Robust Formula
    c_m = ws2.cell(row=row_num, column=13, value=f'=IF(ISNUMBER(K{row_num}), K{row_num}+IF(ISNUMBER(L{row_num}), L{row_num}, 0), "")')
    c_m.alignment = align_right
    c_m.font = font_formula
    c_m.fill = fill_formula
    c_m.number_format = '#,##0'
    
    # N: Selling Price
    c_n = ws2.cell(row=row_num, column=14)
    c_n.alignment = align_right
    c_n.font = font_cell_bold
    c_n.number_format = '#,##0'
    
    # O: Stock Qty
    c_o = ws2.cell(row=row_num, column=15)
    c_o.alignment = align_center
    c_o.font = font_cell_bold
    c_o.number_format = '#,##0'
    
    # P: Inventory Total Value (M * O) - Robust Formula
    c_p = ws2.cell(row=row_num, column=16, value=f'=IF(AND(ISNUMBER(M{row_num}), ISNUMBER(O{row_num})), M{row_num}*O{row_num}, "")')
    c_p.alignment = align_right
    c_p.font = font_formula
    c_p.fill = fill_formula
    c_p.number_format = '#,##0'
    
    # Q: Expected Total Sales (N * O) - Robust Formula
    c_q = ws2.cell(row=row_num, column=17, value=f'=IF(AND(ISNUMBER(N{row_num}), ISNUMBER(O{row_num})), N{row_num}*O{row_num}, "")')
    c_q.alignment = align_right
    c_q.font = font_formula
    c_q.fill = fill_formula
    c_q.number_format = '#,##0'
    
    # R: Margin Rate ((N - M) / M) - Robust Formula
    c_r = ws2.cell(row=row_num, column=18, value=f'=IF(AND(ISNUMBER(M{row_num}), M{row_num}>0, ISNUMBER(N{row_num})), (N{row_num}-M{row_num})/M{row_num}, "")')
    c_r.alignment = align_center
    c_r.font = font_formula
    c_r.fill = fill_formula
    c_r.number_format = '0.0%'
    
    # S: Status
    c_s = ws2.cell(row=row_num, column=19)
    c_s.alignment = align_center
    c_s.font = font_cell_bold
    
    # T: Location
    c_t = ws2.cell(row=row_num, column=20)
    c_t.alignment = align_center
    c_t.font = font_cell
    
    # U: Note
    c_u = ws2.cell(row=row_num, column=21)
    c_u.alignment = align_left
    c_u.font = font_cell

    for c in range(2, 22):
        cell = ws2.cell(row=row_num, column=c)
        cell.border = border_cell
        if cell.fill.fill_type is None and is_alt:
            cell.fill = fill_alt

# -------------------------------------------------------------
# 5. SHEET 3: 사용안내_및_설정 (User Guide & Master Code Tables)
# -------------------------------------------------------------
ws3 = wb.create_sheet(title="사용안내_및_설정")
ws3.views.sheetView[0].showGridLines = True

ws3.column_dimensions['A'].width = 3
ws3.column_dimensions['B'].width = 20
ws3.column_dimensions['C'].width = 18
ws3.column_dimensions['D'].width = 18
ws3.column_dimensions['E'].width = 18
ws3.column_dimensions['F'].width = 3
ws3.column_dimensions['G'].width = 48

# Title
ws3.merge_cells("B2:G2")
ws3["B2"] = "📘 골드랩 엑셀 폼 사용 가이드 및 드롭다운 코드 설정"
ws3["B2"].font = font_title
ws3["B2"].fill = fill_navy_header
ws3["B2"].alignment = Alignment(horizontal="center", vertical="center")
ws3.row_dimensions[2].height = 36

# Guide Box (Right G4:G25)
ws3.merge_cells("G4:G5")
ws3["G4"] = "📸 엑셀에 제품 사진 깔끔하게 넣고 고정하는 핵심 꿀팁"
ws3["G4"].font = font_section
ws3["G4"].fill = fill_section_header
ws3["G4"].alignment = align_left
ws3["G4"].border = border_cell

guide_content = (
    "1. 사진 삽입 방법 (두 가지 중 선택)\n"
    "   • 최신 엑셀: 상단 메뉴 [삽입] ➔ [그림] ➔ [셀에 배치] 클릭 후 사진을 선택하면 해당 셀 크기에 맞게 쏙 들어갑니다.\n"
    "   • 일반 엑셀: [삽입] ➔ [그림] ➔ [이 디바이스]로 사진을 넣고 모서리를 드래그해 셀 크기에 맞춥니다.\n\n"
    "2. 사진이 다른 행으로 밀리거나 엉키지 않게 고정하는 법 (필수!)\n"
    "   • 삽입한 사진을 [마우스 우클릭] ➔ [그림 서식] 클릭\n"
    "   • 우측 패널에서 세 번째 아이콘 [크기 및 속성 (사각형 모양)] 클릭\n"
    "   • '속성' 메뉴에서 【위치와 크기 변동(Move and size with cells)】을 선택하세요!\n"
    "   ➔ 이렇게 하면 행 높이를 조절하거나 필터/정렬을 해도 사진이 셀과 함께 완벽하게 고정됩니다.\n\n"
    "3. 자동 계산 수식 안내 (연노랑 셀)\n"
    "   • 중량(돈): 중량(g)을 입력하면 1돈(3.75g) 기준으로 자동 환산됩니다. (=g/3.75)\n"
    "   • 총 매입원가: 매입단가 + 공임비가 자동으로 합산됩니다.\n"
    "   • 마진율: (판매단가 - 총원가) / 총원가 비율이 실시간 %로 계산됩니다.\n"
    "   • 상단 대시보드: 전체 품목수, 재고수량, 총 재고평가액, 예상 매출액이 실시간 자동 집계됩니다."
)
ws3.merge_cells("G6:G25")
ws3["G6"] = guide_content
ws3["G6"].font = font_cell
ws3["G6"].fill = fill_kpi_bg
ws3["G6"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
for r in range(6, 26):
    ws3.cell(row=r, column=7).border = border_cell

# Master Code Tables (Left B to E)
ws3["B4"] = "카테고리 목록"
ws3["C4"] = "순도/재질 목록"
ws3["D4"] = "판매/진열 상태"
ws3["E4"] = "보관 위치 목록"

for c in range(2, 6):
    cell = ws3.cell(row=4, column=c)
    cell.font = font_th
    cell.fill = fill_navy_header
    cell.alignment = align_center
    cell.border = border_cell

categories = ["순금(24K)", "18K 골드", "14K 골드", "백금(PT)", "은(Silver)", "다이아/보석", "골드바", "코인/메달", "귀금속 원자재", "기타 잡화"]
purities = ["24K (99.99%)", "24K (99.9%)", "18K (75.0%)", "14K (58.5%)", "PT950 (백금)", "Silver 925", "다이아몬드", "합금/기타"]
statuses = ["진열 판매중", "판매 대기", "예약중", "판매 완료", "수리/리셋팅", "금고 보관"]
locations = ["쇼케이스 A-01", "쇼케이스 A-02", "쇼케이스 B-01", "쇼케이스 B-02", "본점 금고 1호실", "본점 금고 2호실", "수탁/위탁 보관함", "지점 이송 대기"]

max_rows = max(len(categories), len(purities), len(statuses), len(locations))
for i in range(max_rows):
    r = 5 + i
    ws3.row_dimensions[r].height = 20
    
    val_b = categories[i] if i < len(categories) else ""
    val_c = purities[i] if i < len(purities) else ""
    val_d = statuses[i] if i < len(statuses) else ""
    val_e = locations[i] if i < len(locations) else ""
    
    for c_idx, val in [(2, val_b), (3, val_c), (4, val_d), (5, val_e)]:
        c_cell = ws3.cell(row=r, column=c_idx, value=val)
        c_cell.font = font_cell
        c_cell.alignment = align_center
        c_cell.border = border_cell
        if i % 2 == 1:
            c_cell.fill = fill_alt

# Add Data Validation Dropdowns to Sheet 2 (목록관리대장)
# Category Dropdown: E8:E30 -> '사용안내_및_설정'!$B$5:$B$14
dv_cat = DataValidation(type="list", formula1="='사용안내_및_설정'!$B$5:$B$14", allow_blank=True)
ws2.add_data_validation(dv_cat)
dv_cat.add("E8:E30")

# Purity Dropdown: G8:G30 -> '사용안내_및_설정'!$C$5:$C$12
dv_purity = DataValidation(type="list", formula1="='사용안내_및_설정'!$C$5:$C$12", allow_blank=True)
ws2.add_data_validation(dv_purity)
dv_purity.add("G8:G30")

# Status Dropdown: S8:S30 -> '사용안내_및_설정'!$D$5:$D$10
dv_status = DataValidation(type="list", formula1="='사용안내_및_설정'!$D$5:$D$10", allow_blank=True)
ws2.add_data_validation(dv_status)
dv_status.add("S8:S30")

# Location Dropdown: T8:T30 -> '사용안내_및_설정'!$E$5:$E$12
dv_loc = DataValidation(type="list", formula1="='사용안내_및_설정'!$E$5:$E$12", allow_blank=True)
ws2.add_data_validation(dv_loc)
dv_loc.add("T8:T30")

# Save Workbook
output_path_workspace = os.path.join(BASE_DIR, "골드랩_제품등록_및_재고관리대장.xlsx")
output_path_downloads = os.path.join(r"C:\Users\WD\Downloads", "골드랩_제품등록_및_재고관리대장.xlsx")

wb.save(output_path_workspace)
wb.save(output_path_downloads)

print(f"Excel file created successfully:")
print(f"Workspace: {output_path_workspace}")
print(f"Downloads: {output_path_downloads}")
